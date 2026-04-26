import os
import re
import json

# import tempfile
from io import BytesIO
from pathlib import Path
from typing import Dict, List

import streamlit as st
from dotenv import load_dotenv
from docx import Document
from openai import OpenAI


# -----------------------------
# Configuration
# -----------------------------
MODEL_NAME = "gpt-4.1"
PLACEHOLDER_PATTERN = re.compile(r"\{\{[A-Z0-9_]+\}\}")


# -----------------------------
# OpenAI setup
# -----------------------------
def get_openai_client() -> OpenAI:
    load_dotenv()
    api_key = os.getenv("OPENAI_API_KEY")

    if not api_key:
        raise ValueError(
            "OPENAI_API_KEY not found. Please create a .env file and add "
            "OPENAI_API_KEY=your_key"
        )

    return OpenAI(api_key=api_key)


# -----------------------------
# DOCX helpers
# -----------------------------
def load_document_from_upload(uploaded_file) -> Document:
    uploaded_file.seek(0)
    return Document(BytesIO(uploaded_file.read()))


def extract_placeholders_from_document(doc: Document) -> List[str]:
    placeholders = set()

    def scan_paragraphs(paragraphs):
        for p in paragraphs:
            for match in PLACEHOLDER_PATTERN.findall(p.text):
                placeholders.add(match)

    scan_paragraphs(doc.paragraphs)

    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                scan_paragraphs(cell.paragraphs)

    for section in doc.sections:
        scan_paragraphs(section.header.paragraphs)
        scan_paragraphs(section.footer.paragraphs)

    return sorted(placeholders)


def replace_placeholders_in_paragraph(paragraph, replacements: Dict[str, str]) -> None:
    original_text = paragraph.text
    new_text = original_text

    for placeholder, value in replacements.items():
        if placeholder in new_text:
            new_text = new_text.replace(placeholder, value)

    if new_text == original_text:
        return

    for run in paragraph.runs:
        run.text = ""

    if paragraph.runs:
        paragraph.runs[0].text = new_text
    else:
        paragraph.add_run(new_text)


def replace_placeholders_in_doc(doc: Document, replacements: Dict[str, str]) -> None:
    for paragraph in doc.paragraphs:
        replace_placeholders_in_paragraph(paragraph, replacements)

    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    replace_placeholders_in_paragraph(paragraph, replacements)

    for section in doc.sections:
        for paragraph in section.header.paragraphs:
            replace_placeholders_in_paragraph(paragraph, replacements)
        for paragraph in section.footer.paragraphs:
            replace_placeholders_in_paragraph(paragraph, replacements)


def document_to_bytes(doc: Document) -> bytes:
    output = BytesIO()
    doc.save(output)
    output.seek(0)
    return output.read()


# -----------------------------
# Code file reading
# -----------------------------
def read_uploaded_code_files(uploaded_files) -> str:
    combined_parts = []

    for uploaded_file in uploaded_files:
        filename = uploaded_file.name

        try:
            uploaded_file.seek(0)
            raw = uploaded_file.read()

            try:
                content = raw.decode("utf-8")
            except UnicodeDecodeError:
                try:
                    content = raw.decode("latin-1")
                except Exception:
                    content = "[Unreadable text file]"
        except Exception as exc:
            content = f"[Error reading file: {exc}]"

        combined_parts.append(
            f"\n\n===== FILE: {filename} =====\nPATH: {filename}\n{content}"
        )

    return "\n".join(combined_parts)


# -----------------------------
# OpenAI generation
# -----------------------------
def build_json_schema(placeholders: List[str]) -> Dict:
    properties = {}
    required = []

    for token in placeholders:
        properties[token] = {
            "type": "string",
            "description": f"English documentation content for placeholder {token}",
        }
        required.append(token)

    return {
        "name": "template_fill",
        "schema": {
            "type": "object",
            "properties": properties,
            "required": required,
            "additionalProperties": False,
        },
        "strict": True,
    }


def generate_placeholder_content(
    client: OpenAI,
    placeholders: List[str],
    code_text: str,
    extra_context: str = "",
) -> Dict[str, str]:
    schema = build_json_schema(placeholders)

    prompt = f"""
You are a senior software documentation engineer.

Your task:
Analyze the provided codebase and generate clear English documentation
for each placeholder.

Rules:
- Return content for every placeholder.
- Be accurate and grounded in the code.
- Do not invent features not present in the code.
- Write in professional English.
- Keep each section useful and readable.
- Use paragraphs and bullet-like text only if it improves clarity.
- If information is missing, explicitly say so briefly.

Optional context from user:
{extra_context if extra_context else "None"}

Placeholders to fill:
{json.dumps(placeholders, indent=2)}

Codebase content:
{code_text[:200000]}
""".strip()

    response = client.responses.create(
        model=MODEL_NAME,
        input=prompt,
        text={
            "format": {
                "type": "json_schema",
                "name": schema["name"],
                "schema": schema["schema"],
                "strict": True,
            }
        },
    )

    raw = response.output_text.strip()
    data = json.loads(raw)

    for token in placeholders:
        if token not in data:
            data[token] = "No content generated."

    return data


# -----------------------------
# UI helpers
# -----------------------------
def get_download_filename(template_name: str) -> str:
    template_path = Path(template_name)
    return f"{template_path.stem}_filled.docx"


def summarize_uploaded_files(uploaded_files) -> str:
    return "\n".join(
        f"{idx}. {f.name}" for idx, f in enumerate(uploaded_files, start=1)
    )


# -----------------------------
# Streamlit App
# -----------------------------
def main() -> None:
    st.set_page_config(page_title="DOCX Template Generator", layout="wide")

    st.title("DOCX Template Documentation Generator")
    st.write(
        "Upload a Word template and one or more code files. "
        "The app will fill placeholders like `{{OVERVIEW}}` in the template."
    )

    with st.sidebar:
        st.header("Configuration")
        st.write(f"Model: `{MODEL_NAME}`")
        st.write("Required environment variable:")
        st.code("OPENAI_API_KEY=your_key", language="bash")

    col1, col2 = st.columns(2)

    with col1:
        st.subheader("1. Upload Word Template")
        template_file = st.file_uploader(
            "Upload template (.docx)",
            type=["docx"],
            accept_multiple_files=False,
        )

        st.subheader("2. Upload Code Files")
        code_files = st.file_uploader(
            "Upload one or more code files",
            accept_multiple_files=True,
            type=[
                "py",
                "js",
                "ts",
                "java",
                "abap",
                "json",
                "xml",
                "yaml",
                "yml",
                "txt",
                "md",
                "sql",
                "html",
                "css",
                "c",
                "cpp",
                "h",
                "hpp",
                "cs",
                "go",
                "rb",
                "php",
                "scala",
                "kt",
                "swift",
            ],
        )

        st.subheader("3. Optional Instructions")
        extra_context = st.text_area(
            "Optional guidance for generated documentation",
            placeholder=(
                "Example:\n"
                "- Write for business users\n"
                "- Emphasize class design\n"
                "- Keep each section short"
            ),
            height=150,
        )

    with col2:
        st.subheader("Selection Summary")

        if template_file:
            st.success(f"Template selected: {template_file.name}")
        else:
            st.info("No template selected.")

        if code_files:
            st.success(f"Code files selected: {len(code_files)}")
            st.text(summarize_uploaded_files(code_files))
        else:
            st.info("No code files selected.")

    if st.button("Generate Document", type="primary"):
        try:
            if not template_file:
                raise ValueError("Please upload a Word template (.docx).")

            if not code_files:
                raise ValueError("Please upload at least one code file.")

            with st.spinner("Reading template..."):
                document = load_document_from_upload(template_file)

            placeholders = extract_placeholders_from_document(document)

            if not placeholders:
                raise ValueError(
                    "No placeholders found in template. "
                    "Use placeholders like {{OVERVIEW}} or {{WORKFLOW}} in the .docx file."
                )

            st.subheader("Detected Placeholders")
            st.code("\n".join(placeholders) if placeholders else "None")

            with st.spinner("Reading code files..."):
                code_text = read_uploaded_code_files(code_files)

            with st.spinner("Generating documentation with OpenAI..."):
                client = get_openai_client()
                replacements = generate_placeholder_content(
                    client=client,
                    placeholders=placeholders,
                    code_text=code_text,
                    extra_context=extra_context,
                )

            with st.spinner("Updating Word document..."):
                replace_placeholders_in_doc(document, replacements)
                output_bytes = document_to_bytes(document)

            download_name = get_download_filename(template_file.name)

            st.success("Document generated successfully.")

            st.subheader("Generated Placeholder Content")
            for key, value in replacements.items():
                with st.expander(key, expanded=False):
                    st.write(value)

            st.download_button(
                label="Download Filled Document",
                data=output_bytes,
                file_name=download_name,
                mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )

        except Exception as exc:
            st.error(f"Error: {exc}")


if __name__ == "__main__":
    main()
